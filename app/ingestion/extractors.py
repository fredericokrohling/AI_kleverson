"""Transforma arquivos (PDF, DOCX, XLSX, PPTX, texto) em texto puro.

Por que texto: é o que a IA lê no bloco 3. O arquivo original continua guardado
(storage.py), então se uma biblioteca melhor aparecer, dá para extrair de novo
(comando `reextrair-anexos`) sem baixar nada do Gmail.

Regra de ouro deste módulo: extrair nunca derruba a coleta. Qualquer falha vira uma
mensagem em português em `error`, que vai para attachments.extraction_error. Assim você
vê exatamente o que não foi lido, em vez de o e-mail inteiro sumir.

Bibliotecas (todas Python puro, sem programas externos no servidor):
- pypdf para PDF;
- python-docx para Word (.docx);
- openpyxl para Excel (.xlsx);
- python-pptx para PowerPoint (.pptx).
"""

# io.BytesIO faz bytes na memória se comportarem como um arquivo aberto.
import io

# dataclass cria classes simples de dados.
from dataclasses import dataclass

# date e datetime aparecem nas células de planilha.
from datetime import date, datetime

# HTMLParser é o leitor de HTML que já vem com o Python.
from html.parser import HTMLParser

# Path pega a extensão do nome do arquivo.
from pathlib import Path


# Resultado de uma extração: ou texto, ou o motivo da falha.
@dataclass(frozen=True)
class Extraction:
    # Texto extraído; None quando não foi possível.
    text: str | None
    # Motivo da falha, em português; None quando deu certo.
    error: str | None = None


# Tipos MIME que identificam cada formato. Os e-mails às vezes mandam o tipo genérico
# "application/octet-stream", por isso a extensão do arquivo também é conferida.
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# Excel moderno.
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# PowerPoint moderno.
_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# Extensões dos formatos antigos do Office (anteriores a 2007), ainda não convertidos.
_LEGACY_OFFICE = {".doc", ".xls", ".ppt"}
# Extensões tratadas como texto simples.
_TEXT_SUFFIXES = {".txt", ".csv", ".tsv", ".md", ".ics", ".json", ".xml", ".log"}


# Descobre o formato pelo nome e pelo tipo MIME. Devolve um rótulo curto.
def detect_kind(filename: str | None, mime_type: str | None) -> str:
    # Extensão em minúsculas (".pdf"); vazio se o nome não tiver extensão.
    suffix = Path(filename or "").suffix.lower()
    # Tipo MIME em minúsculas, sem parâmetros (tira o "; charset=..." do fim).
    mime = (mime_type or "").split(";")[0].strip().lower()
    # PDF.
    if suffix == ".pdf" or mime == "application/pdf":
        return "pdf"
    # Word moderno; .docm é a variante com macros, mesmo formato por dentro.
    if suffix in {".docx", ".docm"} or mime == _DOCX_MIME:
        return "docx"
    # Excel moderno; .xlsm é a variante com macros.
    if suffix in {".xlsx", ".xlsm"} or mime == _XLSX_MIME:
        return "xlsx"
    # PowerPoint moderno.
    if suffix in {".pptx", ".pptm"} or mime == _PPTX_MIME:
        return "pptx"
    # Formatos antigos do Office.
    if suffix in _LEGACY_OFFICE or mime in {
        "application/msword",
        "application/vnd.ms-excel",
        "application/vnd.ms-powerpoint",
    }:
        return "office_antigo"
    # HTML.
    if suffix in {".html", ".htm"} or mime == "text/html":
        return "html"
    # Texto simples (inclui convites de agenda .ics e planilhas .csv).
    if suffix in _TEXT_SUFFIXES or mime.startswith("text/"):
        return "texto"
    # Imagens.
    if mime.startswith("image/"):
        return "imagem"
    # Qualquer outra coisa.
    return "desconhecido"


# Ponto de entrada do módulo: recebe o arquivo e devolve a Extraction.
def extract_text(data: bytes, filename: str | None, mime_type: str | None) -> Extraction:
    """Extrai o texto de um arquivo. Nunca levanta erro: falhas vão em `error`."""
    # Qual formato é.
    kind = detect_kind(filename, mime_type)
    # Dicionário que liga o formato à função que sabe ler.
    readers = {
        # PDF.
        "pdf": _pdf,
        # Word.
        "docx": _docx,
        # Excel.
        "xlsx": _xlsx,
        # PowerPoint.
        "pptx": _pptx,
        # HTML.
        "html": lambda raw: html_to_text(decode_text(raw)),
        # Texto simples.
        "texto": decode_text,
    }
    # Formato antigo do Office: o original fica guardado, mas o texto ainda não sai.
    if kind == "office_antigo":
        # Mensagem explicando o que fazer.
        return Extraction(
            None,
            "Formato antigo do Office (.doc, .xls ou .ppt) ainda não é convertido; "
            "o original está guardado.",
        )
    # Imagem: sem leitura de texto em imagem (OCR) por enquanto.
    if kind == "imagem":
        return Extraction(None, "Imagem: leitura de texto em imagem (OCR) não implementada.")
    # Formato que não sabemos ler.
    if kind not in readers:
        # Informa o tipo para você saber o que apareceu.
        return Extraction(None, f"Tipo de arquivo não suportado: {mime_type or filename}.")
    # Tenta ler; qualquer erro da biblioteca vira mensagem.
    try:
        # Chama a função do formato.
        text = readers[kind](data)
    # Exception pega erros comuns (arquivo corrompido, protegido por senha...).
    except Exception as exc:
        # type(exc).__name__ é o nome do erro (ex.: "BadZipFile"), útil para diagnóstico.
        return Extraction(None, f"Falha ao ler {kind}: {type(exc).__name__}: {exc}")
    # Arquivo lido, mas sem nenhum texto.
    if not text.strip():
        # PDF sem texto quase sempre é um documento escaneado (só imagem).
        if kind == "pdf":
            return Extraction(
                None, "PDF sem texto: provavelmente digitalizado (imagem); OCR não implementado."
            )
        # Outros formatos vazios: grava texto vazio, não é erro.
        return Extraction("")
    # Sucesso.
    return Extraction(text)


# Converte bytes em texto tentando as codificações mais comuns.
def decode_text(data: bytes, charset: str | None = None) -> str:
    # Ordem: a codificação informada pelo e-mail, depois UTF-8, depois a do Windows.
    for encoding in (charset, "utf-8", "cp1252"):
        # Pula a posição vazia (quando não foi informada codificação).
        if not encoding:
            continue
        # Tenta decodificar.
        try:
            # Tenta transformar os bytes em texto com esta codificação.
            return data.decode(encoding)
        # LookupError: codificação desconhecida; UnicodeDecodeError: bytes inválidos.
        except (LookupError, UnicodeDecodeError):
            # Tenta a próxima.
            continue
    # Último recurso: UTF-8 trocando bytes inválidos por "�", sem perder o resto.
    return data.decode("utf-8", errors="replace")


# Leitor de HTML que junta só o texto visível.
class _TextFromHTML(HTMLParser):
    # Tags cujo conteúdo não é texto visível.
    _SKIP = {"script", "style", "head", "title"}
    # Tags que, na tela, começam uma linha nova.
    _BREAK = {"br", "p", "div", "tr", "li", "h1", "h2", "h3", "h4", "h5", "h6", "table"}

    # Prepara o estado do leitor.
    def __init__(self) -> None:
        # convert_charrefs=True transforma "&amp;" em "&" automaticamente.
        super().__init__(convert_charrefs=True)
        # Pedaços de texto acumulados.
        self.parts: list[str] = []
        # Quantas tags "invisíveis" estão abertas agora.
        self._skipping = 0

    # Chamado a cada tag aberta (<p>, <script>...).
    def handle_starttag(self, tag, attrs):
        # Entrou numa tag invisível.
        if tag in self._SKIP:
            # Conta para saber quando sair.
            self._skipping += 1
        # Tag de quebra de linha.
        elif tag in self._BREAK:
            # Acrescenta uma quebra.
            self.parts.append("\n")

    # Chamado a cada tag fechada (</p>...).
    def handle_endtag(self, tag):
        # Saiu de uma tag invisível.
        if tag in self._SKIP and self._skipping:
            # Desconta.
            self._skipping -= 1
        # Fim de parágrafo ou célula de tabela também separa o texto.
        elif tag in {"p", "div", "td", "th", "li"}:
            # Célula vira separador; os outros viram quebra de linha.
            self.parts.append(" | " if tag in {"td", "th"} else "\n")

    # Chamado com o texto entre as tags.
    def handle_data(self, data):
        # Só guarda se não estiver dentro de tag invisível.
        if not self._skipping:
            # Acrescenta o texto.
            self.parts.append(data)


# Converte HTML em texto legível (usado no corpo de e-mails só em HTML).
def html_to_text(html: str) -> str:
    # Cria o leitor.
    parser = _TextFromHTML()
    # Lê o HTML inteiro.
    parser.feed(html)
    # Fecha a leitura (processa o que ficou pendente).
    parser.close()
    # Junta os pedaços e limpa espaços: uma linha por linha não vazia.
    lines = (" ".join(line.split()) for line in "".join(parser.parts).splitlines())
    # Remove linhas vazias e junta com quebra de linha.
    return "\n".join(line for line in lines if line)


# PDF: texto de cada página, com um marcador de página (útil para citar a fonte).
def _pdf(data: bytes) -> str:
    # Import dentro da função: a biblioteca só carrega se aparecer um PDF.
    from pypdf import PdfReader

    # Abre o PDF a partir dos bytes.
    reader = PdfReader(io.BytesIO(data))
    # PDF criptografado: muitos usam senha vazia só para bloquear edição.
    if reader.is_encrypted:
        # decrypt("") devolve 0 quando a senha vazia não abre.
        if not reader.decrypt(""):
            # Levanta erro; extract_text transforma em mensagem.
            raise ValueError("PDF protegido por senha")
    # Lista de páginas com o marcador.
    pages = []
    # enumerate numera a partir de 1.
    for number, page in enumerate(reader.pages, start=1):
        # Texto da página; "or ''" cobre páginas sem texto.
        text = (page.extract_text() or "").strip()
        # Só acrescenta páginas com conteúdo.
        if text:
            # Marcador + texto.
            pages.append(f"### Página {number}\n{text}")
    # Junta as páginas com uma linha em branco entre elas.
    return "\n\n".join(pages)


# Word: parágrafos e tabelas na ordem em que aparecem no documento.
def _docx(data: bytes) -> str:
    # Import local, como no PDF.
    import docx

    # Classe das tabelas, para separar tabela de parágrafo.
    from docx.table import Table

    # Abre o documento.
    document = docx.Document(io.BytesIO(data))
    # Linhas de saída.
    lines = []
    # iter_inner_content percorre parágrafos e tabelas na ordem do documento.
    for block in document.iter_inner_content():
        # Tabela: uma linha de texto por linha da tabela, células separadas por " | ".
        if isinstance(block, Table):
            # Percorre as linhas da tabela.
            for row in block.rows:
                # Texto de cada célula, sem espaços nas pontas.
                cells = [cell.text.strip() for cell in row.cells]
                # Ignora linhas totalmente vazias.
                if any(cells):
                    # Junta as células.
                    lines.append(" | ".join(cells))
        # Parágrafo com texto.
        elif block.text.strip():
            # Acrescenta o parágrafo.
            lines.append(block.text.strip())
    # Uma linha por bloco.
    return "\n".join(lines)


# Formata o valor de uma célula de planilha como texto.
def _cell(value) -> str:
    # Célula vazia.
    if value is None:
        return ""
    # Datas (com ou sem hora) no formato internacional AAAA-MM-DD.
    if isinstance(value, datetime | date):
        # isoformat dá "2026-10-30" ou "2026-10-30T09:00:00".
        return value.isoformat()
    # Qualquer outro valor (texto, número) vira texto.
    return str(value).strip()


# Excel: cada aba com um marcador, linhas com células separadas por tabulação.
# É o caminho dos planos de projeto (PLAN), por isso as datas saem em formato padrão.
def _xlsx(data: bytes) -> str:
    # Import local.
    from openpyxl import load_workbook

    # read_only: lê em fluxo, gasta pouca memória em planilhas grandes.
    # data_only: pega o valor calculado das fórmulas, não a fórmula em si.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    # try/finally garante que o arquivo é fechado mesmo se der erro.
    try:
        # Blocos de saída, um por aba.
        sheets = []
        # Percorre as abas.
        for sheet in workbook.worksheets:
            # Linhas desta aba.
            rows = []
            # values_only devolve só os valores, linha a linha.
            for row in sheet.iter_rows(values_only=True):
                # Formata cada célula.
                cells = [_cell(v) for v in row]
                # Tira as células vazias do fim da linha.
                while cells and not cells[-1]:
                    # Remove a última.
                    cells.pop()
                # Só guarda linhas com algum conteúdo.
                if cells:
                    # Tabulação preserva as colunas.
                    rows.append("\t".join(cells))
            # Aba com conteúdo ganha marcador.
            if rows:
                # Marcador com o nome da aba, depois as linhas.
                sheets.append(f"### Planilha: {sheet.title}\n" + "\n".join(rows))
        # Junta as abas.
        return "\n\n".join(sheets)
    # Sempre executa.
    finally:
        # Fecha o arquivo.
        workbook.close()


# PowerPoint: texto de cada slide (caixas de texto e tabelas) e as anotações.
def _pptx(data: bytes) -> str:
    # Import local.
    from pptx import Presentation

    # Abre a apresentação.
    presentation = Presentation(io.BytesIO(data))
    # Blocos de saída, um por slide.
    slides = []
    # Numera os slides a partir de 1.
    for number, slide in enumerate(presentation.slides, start=1):
        # Linhas deste slide.
        lines = []
        # Percorre as formas do slide (caixas de texto, tabelas, imagens...).
        for shape in slide.shapes:
            # Forma com texto.
            if shape.has_text_frame and shape.text_frame.text.strip():
                # Acrescenta o texto.
                lines.append(shape.text_frame.text.strip())
            # Forma que é tabela.
            if shape.has_table:
                # Percorre as linhas da tabela.
                for row in shape.table.rows:
                    # Junta as células.
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        # Anotações do apresentador, se houver.
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            # Texto das anotações.
            notes = slide.notes_slide.notes_text_frame.text.strip()
            # Só acrescenta se não estiver vazio.
            if notes:
                # Marcadas como anotação.
                lines.append(f"[Anotações] {notes}")
        # Slide com conteúdo ganha marcador.
        if lines:
            # Marcador do slide e o texto.
            slides.append(f"### Slide {number}\n" + "\n".join(lines))
    # Junta os slides.
    return "\n\n".join(slides)
